# -*- coding: utf-8 -*-
# 作用：生成阶段5 中文技术报告（docx，2-4 页）+ 英文摘要，嵌入结果图。
# 运行：& 'D:\eeg-project\.venv\Scripts\python.exe' gen_report.py
import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = r"D:\eeg-project"
FIG = os.path.join(ROOT, "figures")
OUT = os.path.join(ROOT, "results", "step5")
os.makedirs(OUT, exist_ok=True)

doc = Document()


def set_cjk(run, cjk="宋体"):
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), cjk)


# 默认字体
normal = doc.styles["Normal"]
normal.font.name = "Times New Roman"
normal.font.size = Pt(11)
_rpr = normal.element.get_or_add_rPr()
_rpr.get_or_add_rFonts().set(qn("w:eastAsia"), "宋体")


def para(text, bold=False, size=11, align=None, space_after=6):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(size)
    r.font.name = "Times New Roman"
    set_cjk(r, "宋体")
    if align:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    return p


def heading(text, level=1):
    h = doc.add_heading(text, level=level)
    for r in h.runs:
        r.font.name = "Times New Roman"
        set_cjk(r, "黑体")
        r.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
    return h


def figure(fname, caption, width=6.2):
    path = os.path.join(FIG, fname)
    if os.path.exists(path):
        doc.add_picture(path, width=Inches(width))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        c = doc.add_paragraph()
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = c.add_run(caption)
        r.font.size = Pt(9)
        set_cjk(r, "宋体")


# 标题
t = doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run("运动想象脑电分类复现研究：CSP/LDA、跨天迁移与 FBCSP")
r.bold = True; r.font.size = Pt(16); r.font.name = "Times New Roman"; set_cjk(r, "黑体")
sub = doc.add_paragraph(); sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
rs = sub.add_run("作者：张平    日期：2026-10-05    数据集：WBCIC-SHU (2025)")
rs.font.size = Pt(10); rs.font.name = "Times New Roman"; set_cjk(rs, "宋体")

# 英文摘要
heading("Abstract", 1)
para("We reproduced an end-to-end left/right-hand motor imagery (MI) EEG classification pipeline on the "
     "public WBCIC-SHU (2025) dataset (51 subjects x 3 sessions, 58 channels, 250 Hz). Using CSP+LDA with "
     "5-fold stratified cross-validation, within-session accuracy averaged 0.5904 +/- 0.1220. Direct "
     "cross-session transfer dropped to 0.5255 +/- 0.0661, whereas adapting with the first 20% of target-session "
     "trials recovered performance to 0.5718 +/- 0.1091. A filter-bank CSP (FBCSP) pipeline significantly improved "
     "within-session accuracy to 0.6087 +/- 0.1172 (Wilcoxon signed-rank p = 0.0366). All model-fitting steps were "
     "placed inside scikit-learn pipelines to prevent data leakage. Keywords: motor imagery; EEG; CSP; FBCSP; "
     "cross-session transfer; brain-computer interface.")

# 1 背景
heading("1 背景", 1)
para("运动想象脑机接口（MI-BCI）让使用者在没有实际肢体动作的情况下，通过想象肢体运动产生可被解码的脑电信号，"
     "在脑卒中后运动功能康复等神经康复场景中具有重要价值。与其它范式相比，MI 反映使用者的主动意图，但信号信噪比低、"
     "且在不同时间、不同人之间高度不稳定。尤其是同一被试在不同天采集的数据之间（跨 session）存在明显分布漂移，"
     "使得“今天训练的模型明天未必好用”，这是 MI-BCI 走向临床的主要障碍之一。")
para("本研究的目的是：在公开的多日 MI 数据集上，完整复现“预处理—CSP/LDA—跨天迁移—更强的 FBCSP”这一经典流程，"
     "量化跨天性能下降与少量校准的补偿效果，为理解跨天问题与后续研究打基础。")

# 2 数据与方法
heading("2 数据与方法", 1)
heading("2.1 数据", 2)
para("使用上海大学杨帮华团队公开的 WBCIC-SHU (2025) 二分类（2C）数据集：51 名健康被试，每人在不同日期采集 3 次，"
     "每次 200 试次（左手抓握/右手抓握各 100），共 153 个 session，58 导、250 Hz。数据集已由作者完成"
     "去心电/眼电通道、以 Pz 重参考、0.5–40 Hz 带通、去基线、降采样到 250 Hz 等预处理；每个 session 存为一个"
     "维度为 [58 通道 × 1000 采样点 × 200 试次] 的 MAT 文件（采样点对应提示后 4 秒的运动想象段）。")
heading("2.2 分析设置", 2)
para("为避免重复预处理，本文直接使用作者的 processed 数据。分析时仅做一次频段带通与时间窗截取（二者为特征提取、"
     "非“学习”步骤）：频段 8–30 Hz（覆盖 mu 8–12 Hz 与 beta 13–30 Hz），时间窗取运动想象段提示后 0.5–2.5 秒。")
heading("2.3 方法", 2)
para("(1) 同一次采集内：CSP（4 个分量，log 方差）+ 线性判别分析（LDA），5 折分层交叉验证，随机种子固定为 42。")
para("(2) 跨天迁移：用第 1 天全部试次训练，分别在第 2、3 天测试（直接迁移）；适应版训练集 = 第 1 天 + "
     "目标采集前 20% 试次，在目标采集剩余 80% 上测试。")
para("(3) FBCSP：把 4–40 Hz 切为 8 个子带（4–8, 8–12, …, 32–40 Hz），每带各做 CSP（4 分量，log 方差）得到 32 个特征；"
     "用互信息（SelectKBest + mutual_info_classif）选前 8 个，再用 LDA 分类。")
para("防泄漏：CSP、特征选择、分类器全部封装进 scikit-learn Pipeline，在交叉验证内部、仅在训练折上拟合；"
     "跨天实验中训练集与测试集来自不同日期。")
heading("2.4 评估", 2)
para("以分类准确率为指标；FBCSP 与 CSP 的差异用 Wilcoxon 符号秩检验做被试内配对比较。")

# 3 结果
heading("3 结果", 1)
para("信号层面（图 1）：在 C3/C4 可见运动想象相关的节律功率变化，右手想象时对侧 C3 的 mu+beta 功率下降"
     "（相对整段平均 −3.7%）强于同侧 C4（−1.9%），左手想象时对侧 C4（−1.0%）强于同侧 C3（+0.4%），"
     "呈现对侧 ERD（事件相关去同步）的雏形。")
figure("tfr_C3_C4.png", "图 1  sub-001 第 1 次采集：C3/C4 在左右手想象下的时频图（相对整段平均的百分比变化）")

para("同一次采集内（表 1、图 2）：51 名被试平均准确率 0.5904 ± 0.1220；不足 60% 的被试 36 名（70.6%）。"
     "论文报告的 2C 数据 CSP+SVM 平均为 61.12%，本文结果与之接近。")
figure("within_session.png", "图 2  全部被试 within-session CSP+LDA 准确率（按准确率排序）")

para("跨天迁移（表 2、图 3）：同一次采集内平均 0.5904；跨天直接迁移降至 0.5255 ± 0.0661（接近随机水平）；"
     "加入目标采集前 20% 试次适应后回升至 0.5718 ± 0.1091，用少量当次数据即可补回约 4.6 个百分点。")
figure("cross_session.png", "图 3  同一次采集内 vs 跨天直接迁移 vs 跨天+少量适应")

para("FBCSP 对比（表 3、图 4）：FBCSP+LDA 平均 0.6087 ± 0.1172，高于 CSP+LDA 的 0.5904 ± 0.1220；"
     "Wilcoxon 符号秩检验 p = 0.0366，差异具有统计学意义。论文报告的 FBCSP+SVM 为 67.46%。")
figure("fbcsp_vs_csp.png", "图 4  每名被试 CSP vs FBCSP（配对连线），Wilcoxon p = 0.0366")

para("表 1  同一次采集内（CSP+LDA，n=51）", bold=True)
t1 = doc.add_table(rows=1, cols=2); t1.style = "Light Grid Accent 1"
t1.rows[0].cells[0].text = "指标"; t1.rows[0].cells[1].text = "数值（实测）"
for k, v in [("平均准确率", "0.5904 ± 0.1220"), ("低于 60% 的被试", "36 / 51（70.6%）"), ("论文 CSP+SVM 对照", "61.12%")]:
    row = t1.add_row().cells; row[0].text = k; row[1].text = v

para("表 2  跨天迁移（n=51）", bold=True)
t2 = doc.add_table(rows=1, cols=2); t2.style = "Light Grid Accent 1"
t2.rows[0].cells[0].text = "情况"; t2.rows[0].cells[1].text = "平均准确率"
for k, v in [("同一次采集内", "0.5904"), ("跨天直接迁移", "0.5255 ± 0.0661"), ("跨天+前 20% 适应", "0.5718 ± 0.1091")]:
    row = t2.add_row().cells; row[0].text = k; row[1].text = v

para("表 3  FBCSP vs CSP（n=51，配对）", bold=True)
t3 = doc.add_table(rows=1, cols=2); t3.style = "Light Grid Accent 1"
t3.rows[0].cells[0].text = "方法"; t3.rows[0].cells[1].text = "平均准确率"
for k, v in [("CSP + LDA", "0.5904 ± 0.1220"), ("FBCSP + LDA", "0.6087 ± 0.1172"), ("Wilcoxon p", "0.0366（显著）")]:
    row = t3.add_row().cells; row[0].text = k; row[1].text = v

# 4 讨论
heading("4 讨论", 1)
para("跨天差异：跨天直接迁移的性能大幅下降（由 0.59 降到 0.53，接近随机），符合 EEG 非平稳性的预期——"
     "相隔数日后，电极位置、阻抗、被试生理与注意力状态的变化都会造成特征分布漂移，使旧的空间滤波器失配。"
     "而仅用目标采集前 20% 试次重新校准即可显著回补，说明“少量当次校准”对临床康复 BCI 具有实际意义。")
para("与原论文的差异：本文 within-session 平均（0.5904）与论文 CSP+SVM（61.12%）接近；FBCSP（0.6087）"
     "低于论文（67.46%）。差异主要来自方法细节——本文用 4 个 CSP 分量与 LDA，论文用 10 个分量与 SVM；"
     "此外还有时间窗/频段、特征数等设置差异。方向一致：FBCSP 优于 CSP。")
para("局限性：(1) 未采用深度学习，论文用 EEGNet 可达 85.32%，本文未复现；(2) 频段与时间窗为固定选择，"
     "未做个体化调参；(3) 时频图基线采用整段平均（数据无提示前静息段），会低估 ERD 幅度。")
para("【待张平填写：提示——结合你的判断写一段讨论，例如：你如何看待“跨天下降”对临床康复 BCI 落地的影响？"
     "本文结果中“70.6% 的被试低于 60%”是否说明本方法不足，还是部分被试本身难用？请用你自己的话写 3–5 句。】")

# 5 下一步设想
heading("5 下一步设想", 1)
para("【待张平填写：以下是 2–3 个可选方向，供你参考，请勿照抄，结合你的中医脑病临床背景选定或改写：")
para("  方向 A（临床落地）：面向卒中后偏瘫手功能康复，研究“每日少量校准”方案在真实康复训练中的可行性，"
     "比较不同校准试次数对当天 BCI 可用性的影响。", space_after=2)
para("  方向 B（方法改进）：引入迁移学习/域自适应或深度学习（EEGNet/FBCNet）来缓解跨天漂移，"
     "并与传统 CSP/FBCSP 做系统对比。", space_after=2)
para("  方向 C（临床人群）：把研究扩展到脑卒中、意识障碍等病人群体，考察其运动想象能力与 ERD 特征，"
     "探索 BCI 用于康复评估与促醒的可能。】", space_after=6)

# 6 参考文献
heading("6 参考文献", 1)
para("[1] Yang, B., Rong, F., Xie, Y., Li, D., Zhang, J., Li, F., Shi, G., & Gao, X. (2025). "
     "A multi-day and high-quality EEG dataset for motor imagery brain-computer interface. "
     "Scientific Data, 12, 488. DOI: 10.1038/s41597-025-04826-y.", size=10)
para("[2] Ang, K. K., Chin, Z. Y., Zhang, H., & Guan, C. (2008). Filter bank common spatial pattern (FBCSP) "
     "in brain-computer interface. IEEE IJCNN, 2390–2397.", size=10)
para("[3] Lawhern, V. J., et al. (2018). EEGNet: a compact convolutional neural network for EEG-based "
     "brain-computer interfaces. Journal of Neural Engineering, 15(5), 056013.", size=10)
para("[4] Pfurtscheller, G., & Lopes da Silva, F. H. (1999). Event-related EEG/MEG synchronization and "
     "desynchronization: basic principles. Clinical Neurophysiology, 110(11), 1842–1857.", size=10)
para("[5] 张平，本项目 GitHub 仓库：https://github.com/jumping158/wbcic-shu-motor-imagery-eeg", size=10)
para("说明：本文代码在 AI 辅助下完成，分析与结果解释由作者负责。", size=10)

path = os.path.join(OUT, "技术报告.docx")
doc.save(path)
print("已写", path)
